import streamlit as st
import pandas as pd
import io
import os
from datetime import date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Modul untuk pembuatan dokumen Word (.docx) berstandar instrumen resmi
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Konfigurasi Halaman
st.set_page_config(
    page_title="Monev TKA - Cabdin Bangkalan",
    page_icon="📋",
    layout="wide"
)

# KATA SANDI RESET PERMANEN
PASSWORD_ADMIN = "admin123"

# Nama berkas Excel permanen di server/komputer
MASTER_FILE = "master_data_monev_tka_2026.xlsx"

# 19 Komponen Instrumen Monev TKA Berdasarkan Dokumen Resmi
komponen_tka = [
    "Satuan Pendidikan Pelaksana TKA telah menetapkan pembagian sesi untuk setiap peserta beserta komputer yang akan digunakan selama TKA",
    "Ruang TKA, Perangkat Komputer, Kartu Login Peserta dan Daftar Hadir sudah dipersiapkan paling lambat 1 (satu) hari sebelum pelaksanaan TKA",
    "Jadwal TKA tertulis rinci dan jelas",
    "Setiap ruang TKA disediakan denah tempat duduk peserta",
    "Setiap Meja/PC dalam ruang ujian diberi nomor peserta TKA",
    "Setiap Peserta memiliki kartu peserta TKA",
    "Terdapat tata tertib pengawas ruang, Proktor dan Teknisi",
    "Tata tertib peserta TKA ditempel pada dinding",
    "Setiap Ruangan diawasi oleh Pengawas Silang",
    "Semua Pengawas dan Panitia TKA mengisi Daftar Hadir",
    "Semua Peserta TKA mengisi Daftar hadir",
    "Berita Acara tersedia dan diisi",
    "Tiap ruang asesmen ditempelkan pengumuman yang bertuliskan 'SELAIN PESERTA, PENGAWAS, PROKTOR, DAN TEKNISI DILARANG MASUK RUANG TES'",
    "Di lokasi asesmen dipasang pengumuman 'SEDANG BERLANGSUNG TKA'",
    "Tiap ruang asesmen ditempelkan pengumuman yang bertuliskan 'DILARANG MEMBAWA PERANGKAT KOMUNIKASI ELEKTRONIK, KAMERA, DAN SEJENISNYA KE DALAM RUANG TKA'",
    "Perangkat komunikasi elektronik, kamera dan sejenisnya tidak berada di Ruang TKA",
    "Setiap ruang TKA memiliki Pencahayaan dan Ventilasi/ Sirkulasi Udara yang cukup",
    "Pengawas Ruang, Proktor, dan Teknisi telah menandatangani Fakta Integritas",
    "Keberadaan Kepala Sekolah"
]

# Fungsi Format Tanggal ke Bahasa Indonesia (Contoh: 22 Oktober 2026)
def format_indo_date(date_str):
    months = {
        '01': 'Januari', '02': 'Februari', '03': 'Maret', '04': 'April',
        '05': 'Mei', '06': 'Juni', '07': 'Juli', '08': 'Agustus',
        '09': 'September', '10': 'Oktober', '11': 'November', '12': 'Desember'
    }
    try:
        parts = str(date_str).split('-')
        if len(parts) == 3:
            y, m, d = parts
            return f"{int(d)} {months.get(m, m)} {y}"
    except:
        pass
    return date_str

# Fungsi Membaca Data Master dari Berkas Excel Permanen
def load_master_data():
    if os.path.exists(MASTER_FILE):
        try:
            xl = pd.ExcelFile(MASTER_FILE)
            if "SMK" in xl.sheet_names:
                df_smk = pd.read_excel(MASTER_FILE, sheet_name="SMK", skiprows=2)
            else:
                df_smk = pd.DataFrame()
                
            if "SMA" in xl.sheet_names:
                df_sma = pd.read_excel(MASTER_FILE, sheet_name="SMA", skiprows=2)
            else:
                df_sma = pd.DataFrame()
                
            df_combined = pd.concat([df_smk, df_sma], ignore_index=True)
            if "No" in df_combined.columns:
                df_combined = df_combined.drop(columns=["No"])
            return df_combined.drop_duplicates(subset=["Satuan Pendidikan"])
        except Exception:
            try:
                df = pd.read_excel(MASTER_FILE, skiprows=2)
                if "No" in df.columns:
                    df = df.drop(columns=["No"])
                return df
            except:
                return pd.DataFrame()
    return pd.DataFrame()

# Helper untuk membuat struktur sheet Excel dengan Judul & Auto-Border
def create_styled_sheet(ws, sheet_title, df_subset):
    ws.title = sheet_title
    ws.views.sheetView[0].showGridLines = True

    PRIMARY_COLOR = "1F4E78"
    font_title = Font(name="Calibri", size=13, bold=True, color=PRIMARY_COLOR)
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10)
    
    fill_header = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin_side = Side(style="thin", color="D9D9D9")
    border_thin = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    border_header = Border(left=thin_side, right=thin_side, top=thin_side, bottom=Side(style="medium", color=PRIMARY_COLOR))

    if not df_subset.empty:
        df_save = df_subset.copy()
        if "No" in df_save.columns:
            df_save = df_save.drop(columns=["No"])
        df_save.insert(0, "No", range(1, len(df_save) + 1))
    else:
        dummy_cols = ["No", "Satuan Pendidikan", "Nama Petugas", "Tanggal Observasi", "Status Pelaksanaan", "Moda Pelaksanaan", "Sesi 1 (Siswa)", "Sesi 2 (Siswa)", "Sesi 3 (Siswa)", "Sesi 4 (Siswa)", "Jumlah Ruang", "Jumlah Pengawas", "Jumlah Proktor", "Rekomendasi Petugas"]
        df_save = pd.DataFrame(columns=dummy_cols)

    headers = list(df_save.columns)
    max_col_idx = len(headers)
    last_col_letter = get_column_letter(max_col_idx)

    ws.merge_cells(f"A1:{last_col_letter}1")
    cell_title = ws["A1"]
    cell_title.value = f"Rekapitulasi hasil Monitoring dan Evaluasi Pelaksanaan TKA Tahun 2026 Jenjang {sheet_title} Cabang Dinas Pendidikan Wilayah Kabupaten Bangkalan"
    cell_title.font = font_title
    cell_title.alignment = align_left
    ws.row_dimensions[1].height = 32

    ws.row_dimensions[2].height = 12

    ws.append([])
    ws.append(headers)
    ws.row_dimensions[3].height = 38

    for col_n in range(1, max_col_idx + 1):
        cell = ws.cell(row=3, column=col_n)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_header

    actual_rows_count = len(df_save)
    for i, r_val in enumerate(df_save.values, start=4):
        ws.append(list(r_val))
        ws.row_dimensions[i].height = 26
        for col_n in range(1, len(r_val) + 1):
            cell = ws.cell(row=i, column=col_n)
            cell.font = font_data
            if col_n in [1, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13]:
                cell.alignment = align_center
            elif col_n in [2, 3, 14]:
                cell.alignment = align_left
            else:
                cell.alignment = align_center
            cell.border = border_thin

    start_empty_row = 4 + actual_rows_count
    total_rows_to_render = max(15, 5 - actual_rows_count)
    for empty_row_idx in range(start_empty_row, start_empty_row + total_rows_to_render):
        ws.row_dimensions[empty_row_idx].height = 22
        for col_n in range(1, max_col_idx + 1):
            cell = ws.cell(row=empty_row_idx, column=col_n)
            cell.border = border_thin
            if col_n == 1:
                cell.value = ""

    for col_idx in range(1, max_col_idx + 1):
        col_letter = get_column_letter(col_idx)
        if col_idx == 1:
            ws.column_dimensions[col_letter].width = 5
        elif col_idx in [2, 3]:
            ws.column_dimensions[col_letter].width = 22
        elif col_idx == 4:
            ws.column_dimensions[col_letter].width = 13
        elif col_idx in [5, 6]:
            ws.column_dimensions[col_letter].width = 12
        elif col_idx in [7, 8, 9, 10]:
            ws.column_dimensions[col_letter].width = 10
        elif col_idx in [11, 12, 13]:
            ws.column_dimensions[col_letter].width = 11
        elif col_idx == 14:
            ws.column_dimensions[col_letter].width = 25
        else:
            ws.column_dimensions[col_letter].width = 15

# Fungsi Menyimpan Data ke Master Excel dengan 2 Sheet Terpisah (SMA & SMK)
def save_to_master_excel(row_data):
    df_existing = load_master_data()
    
    if not df_existing.empty and "Satuan Pendidikan" in df_existing.columns:
        idx = df_existing[df_existing["Satuan Pendidikan"] == row_data["Satuan Pendidikan"]].index
        if not idx.empty:
            df_existing.loc[idx[0]] = row_data
        else:
            df_existing = pd.concat([df_existing, pd.DataFrame([row_data])], ignore_index=True)
    else:
        df_existing = pd.DataFrame([row_data])
        
    wb = openpyxl.Workbook()
    
    if "Satuan Pendidikan" in df_existing.columns:
        df_smk = df_existing[df_existing["Satuan Pendidikan"].astype(str).str.contains("SMK", case=False, na=False)]
        df_sma = df_existing[df_existing["Satuan Pendidikan"].astype(str).str.contains("SMA", case=False, na=False)]
    else:
        df_smk = pd.DataFrame()
        df_sma = pd.DataFrame()

    ws_smk = wb.active
    create_styled_sheet(ws_smk, "SMK", df_smk)

    ws_sma = wb.create_sheet(title="SMA")
    create_styled_sheet(ws_sma, "SMA", df_sma)

    wb.save(MASTER_FILE)
    return df_existing

# Helper untuk mewarnai sel tabel Word
def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

# Helper untuk memberi border tipis rapi pada tabel Word
def set_table_borders(table, color="D3D3D3", sz="4"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

# Fungsi Format Resmi Lembar Instrumen Word (.docx) dengan Kotak Rekomendasi & Tanggal Indonesia
def generate_docx_report(sch_data, komponen_list):
    doc = docx.Document()
    
    # Atur Margin Halaman (0.75 inci di kiri & kanan)
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Gaya Font Default
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)

    # 1. JUDUL INSTRUMEN RESMI (DIPERBESAR & SIMETRIS)
    p_h1 = doc.add_paragraph()
    p_h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_h1.paragraph_format.space_after = Pt(2)
    r1 = p_h1.add_run("INSTRUMEN MONITORING DAN EVALUASI PELAKSANAAN TKA\n")
    r1.bold = True
    r1.font.size = Pt(13.5)

    r2 = p_h1.add_run("CABANG DINAS PENDIDIKAN WILAYAH KABUPATEN BANGKALAN TAHUN 2026")
    r2.bold = True
    r2.font.size = Pt(11.5)

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(10)
    r_line = p_line.add_run("―" * 46)
    r_line.bold = True
    r_line.font.color.rgb = RGBColor(80, 80, 80)

    # 2. INFORMASI IDENTITAS SEKOLAH (TITIK DUA RAPI SEMPURNA)
    identities = [
        ("Hari / Tanggal", format_indo_date(sch_data.get('Tanggal Observasi', '-'))),
        ("Nama Sekolah Penyelenggara", sch_data.get('Satuan Pendidikan', '-')),
        ("Status Pelaksanaan", f"{sch_data.get('Status Pelaksanaan', 'Mandiri')}"),
        ("Moda Pelaksanaan", f"{sch_data.get('Moda Pelaksanaan', 'Online')}"),
        ("Nama Petugas / Observer", sch_data.get('Nama Petugas', '-'))
    ]

    tbl_id = doc.add_table(rows=len(identities), cols=2)
    tbl_id.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_id.autofit = False
    set_table_borders(tbl_id, color="FFFFFF", sz="0")

    for idx, (label, val) in enumerate(identities):
        cell_lbl = tbl_id.cell(idx, 0)
        p_l = cell_lbl.paragraphs[0]
        p_l.paragraph_format.space_after = Pt(3)
        r_l = p_l.add_run(label)
        r_l.bold = True

        cell_val = tbl_id.cell(idx, 1)
        p_v = cell_val.paragraphs[0]
        p_v.paragraph_format.space_after = Pt(3)
        r_v = p_v.add_run(f": {val}")
        if label in ["Nama Sekolah Penyelenggara", "Nama Petugas / Observer"]:
            r_v.bold = True

    tbl_id.rows[0].cells[0].width = Inches(2.3)
    tbl_id.rows[0].cells[1].width = Inches(4.7)

    # 3. TABEL REKAPITULASI JUMLAH SISWA, SESI, DAN RUANG
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(8)
    p_space.paragraph_format.space_after = Pt(4)
    r_sub1 = p_space.add_run("Jumlah Siswa TKA, Sesi & Ruang Ujian:")
    r_sub1.bold = True

    tbl_s = doc.add_table(rows=2, cols=7)
    tbl_s.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_s.autofit = False
    set_table_borders(tbl_s, color="B0C4DE", sz="4")

    headers_s = ["Sesi 1", "Sesi 2", "Sesi 3", "Sesi 4", "Jml Ruang", "Pengawas", "Proktor"]
    col_widths_s = [Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0), Inches(1.0)]

    vals_s = [
        str(sch_data.get('Sesi 1 (Siswa)', 0)),
        str(sch_data.get('Sesi 2 (Siswa)', 0)),
        str(sch_data.get('Sesi 3 (Siswa)', 0)),
        str(sch_data.get('Sesi 4 (Siswa)', 0)),
        str(sch_data.get('Jumlah Ruang', 1)),
        str(sch_data.get('Jumlah Pengawas', 1)),
        str(sch_data.get('Jumlah Proktor', 1))
    ]

    for col_idx, h_text in enumerate(headers_s):
        cell = tbl_s.cell(0, col_idx)
        set_cell_background(cell, "333333")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    for col_idx, v_text in enumerate(vals_s):
        cell = tbl_s.cell(1, col_idx)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(v_text)
        r.bold = True
        r.font.size = Pt(10)

    for row in tbl_s.rows:
        for idx, width in enumerate(col_widths_s):
            row.cells[idx].width = width

    # 4. TABEL 19 KOMPONEN YANG DIMONITOR (DENGAN REPEAT HEADER DI HALAMAN BERIKUTNYA)
    p_space2 = doc.add_paragraph()
    p_space2.paragraph_format.space_before = Pt(12)
    p_space2.paragraph_format.space_after = Pt(4)
    r_sub2 = p_space2.add_run("Komponen Yang Dimonitor:")
    r_sub2.bold = True

    tbl_m = doc.add_table(rows=len(komponen_list) + 1, cols=5)
    tbl_m.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_m.autofit = False
    set_table_borders(tbl_m, color="D3D3D3", sz="4")

    headers_m = ["No", "Komponen Yang Dimonitor", "Ada dan Lengkap", "Ada tetapi Tidak Lengkap", "Tidak Ada"]
    col_widths_m = [Inches(0.5), Inches(3.5), Inches(1.0), Inches(1.1), Inches(0.9)]

    header_tr = tbl_m.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    header_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

    for c_idx, h_text in enumerate(headers_m):
        cell = tbl_m.cell(0, c_idx)
        set_cell_background(cell, "333333")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    for r_idx, k_text in enumerate(komponen_list, start=1):
        status_ans = sch_data.get(f"Komponen {r_idx}", "Ada dan Lengkap")
        
        row_tr = tbl_m.rows[r_idx]._tr.get_or_add_trPr()
        row_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        c_no = tbl_m.cell(r_idx, 0)
        p_no = c_no.paragraphs[0]
        p_no.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_no.add_run(str(r_idx)).bold = True

        c_k = tbl_m.cell(r_idx, 1)
        p_k = c_k.paragraphs[0]
        p_k.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_k.add_run(k_text)

        for col_ans_idx, col_name in enumerate(["Ada dan Lengkap", "Ada tetapi Tidak Lengkap", "Tidak Ada"], start=2):
            c_ans = tbl_m.cell(r_idx, col_ans_idx)
            p_ans = c_ans.paragraphs[0]
            p_ans.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if status_ans == col_name:
                r_chk = p_ans.add_run("✓")
                r_chk.bold = True
                r_chk.font.size = Pt(12)
                r_chk.font.color.rgb = RGBColor(51, 51, 51)

    for row in tbl_m.rows:
        for idx, width in enumerate(col_widths_m):
            row.cells[idx].width = width

    # 5. HASIL KUNJUNGAN DAN REKOMENDASI PETUGAS (DIBERIKAN KOTAK / BOX KHUSUS)
    p_rek_head = doc.add_paragraph()
    p_rek_head.paragraph_format.space_before = Pt(12)
    p_rek_head.paragraph_format.space_after = Pt(4)
    r_rh = p_rek_head.add_run("HASIL KUNJUNGAN DAN REKOMENDASI PETUGAS :")
    r_rh.bold = True

    # Tabel Kotak untuk Catatan Rekomendasi
    tbl_box = doc.add_table(rows=1, cols=1)
    tbl_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_box.autofit = False
    tbl_box.rows[0].cells[0].width = Inches(7.0)
    set_table_borders(tbl_box, color="333333", sz="6") # Border kotak tegas rapi
    
    cell_box = tbl_box.cell(0, 0)
    p_box = cell_box.paragraphs[0]
    p_box.paragraph_format.space_before = Pt(6)
    p_box.paragraph_format.space_after = Pt(6)
    p_box.paragraph_format.left_indent = Inches(0.1)
    p_box.paragraph_format.right_indent = Inches(0.1)
    
    rekomendasi_teks = str(sch_data.get('Rekomendasi Petugas', '-'))
    p_box.add_run(rekomendasi_teks)

    # Spasi sebelum tanda tangan
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(10)

    # 6. TANDA TANGAN BERDAMPINGAN DENGAN TANGGAL INDONESIA (SIMETRIS SEMPURNA)
    tbl_ttd = doc.add_table(rows=4, cols=2)
    tbl_ttd.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ttd.autofit = False
    set_table_borders(tbl_ttd, color="FFFFFF", sz="0")

    tgl_indo = format_indo_date(sch_data.get('Tanggal Observasi', str(date.today())))
    p_d = tbl_ttd.cell(0, 1).paragraphs[0]
    p_d.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_d.add_run(f"Bangkalan, {tgl_indo}").bold = True

    p00 = tbl_ttd.cell(1, 0).paragraphs[0]
    p00.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p00.add_run("Mengetahui,\nKepala Sekolah / Mewakili").bold = True

    p01 = tbl_ttd.cell(1, 1).paragraphs[0]
    p01.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p01.add_run("Petugas Monev (Observer)").bold = True

    tbl_ttd.cell(2, 0).paragraphs[0].add_run("\n\n\n")
    tbl_ttd.cell(2, 1).paragraphs[0].add_run("\n\n\n")

    p20 = tbl_ttd.cell(3, 0).paragraphs[0]
    p20.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p20.add_run("( ___________________________ )").bold = True

    p21 = tbl_ttd.cell(3, 1).paragraphs[0]
    p21.alignment = WD_ALIGN_PARAGRAPH.CENTER
    petugas_nama = sch_data.get('Nama Petugas', 'Yustinus Budi Setyanta')
    p21.add_run(f"( {petugas_nama} )").bold = True

    for row in tbl_ttd.rows:
        row.cells[0].width = Inches(3.5)
        row.cells[1].width = Inches(3.5)

    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

# --- INTERFACE STREAMLIT ---
st.title("📋 LEMBAR OBSERVASI MONEV PELAKSANAAN TKA")

st.markdown(
    """
    <p style='color: #1F4E78; font-size: 1.05rem; font-weight: 600; margin-top: -10px; margin-bottom: 20px;'>
        Jenjang SMA dan SMK Cabang Dinas Pendidikan Wilayah Kabupaten Bangkalan Tahun 2026
    </p>
    """,
    unsafe_allow_html=True
)
st.markdown("---")

df_master_current = load_master_data()

# --- SIDEBAR PANEL KONTROL ---
st.sidebar.header("⚙️ Panel Kontrol Master")
st.sidebar.info(f"Total Sekolah Tersimpan: **{len(df_master_current)}**")

edit_mode = False
selected_school_data = None

if not df_master_current.empty:
    st.sidebar.markdown("---")
    st.sidebar.subheader("✏️ Mode Edit Data")
    list_sekolah = df_master_current["Satuan Pendidikan"].tolist()
    pilihan_edit = st.sidebar.selectbox("Pilih Sekolah untuk Diedit:", ["-- Input Baru --"] + list_sekolah)
    
    if pilihan_edit != "-- Input Baru --":
        edit_mode = True
        selected_school_data = df_master_current[df_master_current["Satuan Pendidikan"] == pilihan_edit].iloc[0]

with st.sidebar.expander("🔐 Fitur Reset / Hapus Database", expanded=False):
    st.caption("Memerlukan Kata Sandi Administrator")
    pwd_input = st.text_input("Masukkan Password Admin:", type="password", key="pwd_reset")
    
    if st.button("🗑️ Konfirmasi Hapus Permanen", type="secondary"):
        if pwd_input == PASSWORD_ADMIN:
            if os.path.exists(MASTER_FILE):
                os.remove(MASTER_FILE)
                st.success("✅ Seluruh Database berhasil dikosongkan!")
                st.rerun()
            else:
                st.info("Database sudah dalam keadaan kosong.")
        else:
            st.error("❌ Kata sandi salah! Akses ditolak.")

if edit_mode:
    st.warning(f"⚠️ **Anda sedang dalam MODE EDIT** untuk sekolah: **{pilihan_edit}**. Mengklik simpan akan memperbarui data sekolah ini.")

# --- FORM IDENTITAS & PELAKSANAAN SEKOLAH ---
st.subheader("📌 A. Informasi Umum & Pelaksanaan")
c1, c2 = st.columns(2)

default_petugas = selected_school_data["Nama Petugas"] if edit_mode else "Yustinus Budi Setyanta"
default_sekolah = selected_school_data["Satuan Pendidikan"] if edit_mode else "SMK Negeri 1 Bangkalan"
default_tanggal = pd.to_datetime(selected_school_data["Tanggal Observasi"]).date() if edit_mode else date.today()
default_status = selected_school_data["Status Pelaksanaan"] if edit_mode else "Mandiri"
default_moda = selected_school_data["Moda Pelaksanaan"] if edit_mode else "Online"

with c1:
    satuan_pendidikan = st.text_input("Nama Sekolah Penyelenggara", value=default_sekolah, disabled=edit_mode)
    if edit_mode:
        st.caption("ℹ️ Nama sekolah terkunci saat mode edit.")
    nama_petugas = st.text_input("Nama Petugas / Observer", value=default_petugas)
    tanggal_observasi = st.date_input("Hari / Tanggal Observasi", value=default_tanggal)

with c2:
    status_pelaksanaan = st.selectbox("Status Pelaksanaan", ["Mandiri", "Menumpang"], index=0 if default_status=="Mandiri" else 1)
    moda_pelaksanaan = st.selectbox("Moda Pelaksanaan", ["Online", "Semi Online"], index=0 if default_moda=="Online" else 1)

# --- REKAPITULASI JUMLAH SISWA TKA ---
st.subheader("📊 B. Rekapitulasi Jumlah Siswa TKA, Sesi & Ruang")

st.markdown(
    """
    <div style="background-color: #dbeafe; padding: 12px; border-radius: 6px; border-left: 5px solid #1F4E78; color: #1e3a8a; font-weight: 600;">
        ℹ️ Masukkan jumlah siswa, pengawas, proktor, dan ruang sesuai pelaksanaan di sekolah.
    </div>
    """,
    unsafe_allow_html=True
)
st.markdown("")

col_s1, col_s2, col_s3, col_s4, col_s5 = st.columns(5)
with col_s1:
    jml_sesi_1 = st.number_input("Sesi 1", min_value=0, value=int(selected_school_data["Sesi 1 (Siswa)"]) if edit_mode and "Sesi 1 (Siswa)" in selected_school_data else 0)
with col_s2:
    jml_sesi_2 = st.number_input("Sesi 2", min_value=0, value=int(selected_school_data["Sesi 2 (Siswa)"]) if edit_mode and "Sesi 2 (Siswa)" in selected_school_data else 0)
with col_s3:
    jml_sesi_3 = st.number_input("Sesi 3", min_value=0, value=int(selected_school_data["Sesi 3 (Siswa)"]) if edit_mode and "Sesi 3 (Siswa)" in selected_school_data else 0)
with col_s4:
    jml_sesi_4 = st.number_input("Sesi 4", min_value=0, value=int(selected_school_data["Sesi 4 (Siswa)"]) if edit_mode and "Sesi 4 (Siswa)" in selected_school_data else 0)
with col_s5:
    jml_ruang = st.number_input("Jumlah Ruang Ujian", min_value=1, value=int(selected_school_data["Jumlah Ruang"]) if edit_mode and "Jumlah Ruang" in selected_school_data else 1)

col_p1, col_p2 = st.columns(2)
with col_p1:
    jml_pengawas = st.number_input("Jumlah Pengawas Ruang", min_value=0, value=int(selected_school_data["Jumlah Pengawas"]) if edit_mode and "Jumlah Pengawas" in selected_school_data else 1)
with col_p2:
    jml_proktor = st.number_input("Jumlah Proktor", min_value=0, value=int(selected_school_data["Jumlah Proktor"]) if edit_mode and "Jumlah Proktor" in selected_school_data else 1)

st.markdown("---")
st.subheader("📝 C. Komponen Yang Dimonitor (19 Butir Instrumen)")

opsi_jawaban = ["Ada dan Lengkap", "Ada tetapi Tidak Lengkap", "Tidak Ada"]
jawaban = {}

for i, kriteria in enumerate(komponen_tka, start=1):
    default_ans_idx = 0
    key_name = f"komponen_{i}"
    
    if edit_mode and key_name in selected_school_data:
        val_lama = selected_school_data[key_name]
        if val_lama == "Ada tetapi Tidak Lengkap":
            default_ans_idx = 1
        elif val_lama == "Tidak Ada":
            default_ans_idx = 2
            
    with st.container():
        st.markdown(f"**{i}. {kriteria}**")
        ans = st.radio(
            label=f"Pilih status untuk nomor {i}",
            options=opsi_jawaban,
            index=default_ans_idx,
            horizontal=True,
            key=key_name,
            label_visibility="collapsed"
        )
        jawaban[key_name] = ans
        st.markdown("")

st.markdown("---")
st.subheader("📌 D. Hasil Kunjungan dan Rekomendasi Petugas")
default_rekomendasi = selected_school_data["Rekomendasi Petugas"] if edit_mode and "Rekomendasi Petugas" in selected_school_data else ""
rekomendasi_petugas = st.text_area("Masukkan Catatan Hasil Kunjungan dan Rekomendasi:", value=default_rekomendasi)

st.markdown("---")

btn_label = "💾 Perbarui Data Sekolah (Edit)" if edit_mode else "💾 Simpan Permanen ke Master Excel"
if st.button(btn_label, type="primary"):
    if not (nama_petugas and satuan_pendidikan):
        st.warning("⚠️ Mohon lengkapi Nama Petugas dan Satuan Pendidikan terlebih dahulu!")
    else:
        row_data = {
            "Satuan Pendidikan": satuan_pendidikan,
            "Nama Petugas": nama_petugas,
            "Tanggal Observasi": str(tanggal_observasi),
            "Status Pelaksanaan": status_pelaksanaan,
            "Moda Pelaksanaan": moda_pelaksanaan,
            "Sesi 1 (Siswa)": jml_sesi_1,
            "Sesi 2 (Siswa)": jml_sesi_2,
            "Sesi 3 (Siswa)": jml_sesi_3,
            "Sesi 4 (Siswa)": jml_sesi_4,
            "Jumlah Ruang": jml_ruang,
            "Jumlah Pengawas": jml_pengawas,
            "Jumlah Proktor": jml_proktor,
            "Rekomendasi Petugas": rekomendasi_petugas
        }
        
        for idx, (k_key, k_val) in enumerate(jawaban.items(), start=1):
            row_data[f"Komponen {idx}"] = k_val
        
        save_to_master_excel(row_data)
        if edit_mode:
            st.success(f"✅ Data untuk **{satuan_pendidikan}** berhasil diperbarui!")
        else:
            st.success(f"✅ Data untuk **{satuan_pendidikan}** berhasil disimpan secara PERMANEN ke database Master!")
        st.rerun()

# --- FITUR UNDUH LAPORAN INDIVIDU SEKOLAH (.DOCX) ---
df_display = load_master_data()
if not df_display.empty:
    st.markdown("---")
    st.subheader("🖨️ Cetak / Unduh Lembar Monev Per Sekolah (.docx)")
    
    pilih_sekolah_cetak = st.selectbox("Pilih Satuan Pendidikan untuk Dicetak Laporannya:", df_display["Satuan Pendidikan"].tolist())
    
    if pilih_sekolah_cetak:
        sch_data = df_display[df_display["Satuan Pendidikan"] == pilih_sekolah_cetak].iloc[0]
        docx_bytes = generate_docx_report(sch_data, komponen_tka)

        st.download_button(
            label=f"📄 Unduh Dokumen Word Laporan ({pilih_sekolah_cetak})",
            data=docx_bytes,
            file_name=f"Laporan_Monev_{pilih_sekolah_cetak.replace(' ', '_')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary"
        )

# --- TABEL MASTER REKAP PERMANEN ---
if not df_display.empty:
    st.markdown("---")
    st.subheader("📊 Tabel Master Rekap Permanen (Semua Sekolah)")
    
    df_show = df_display.copy()
    df_show.insert(0, "No", range(1, len(df_show) + 1))
    st.dataframe(df_show, hide_index=True)

    with open(MASTER_FILE, "rb") as f:
        excel_bytes = f.read()

    st.download_button(
        label="📥 Unduh Master Excel (.xlsx) - Semua Sekolah",
        data=excel_bytes,
        file_name="Master_Monev_TKA_Cabdin_Bangkalan.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="secondary"
    )